#!/usr/bin/env python3
import sys
from docx import Document

def extract_docx_text(docx_path, output_path):
    doc = Document(docx_path)
    
    with open(output_path, 'w', encoding='utf-8') as f:
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                f.write(text + '\n\n')
        
        # Extract tables
        for table in doc.tables:
            f.write('\n--- TABLE ---\n')
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    row_text.append(cell.text.strip())
                f.write(' | '.join(row_text) + '\n')
            f.write('--- END TABLE ---\n\n')

if __name__ == '__main__':
    if len(sys.argv) != 3:
        print("Usage: python extract_docx.py <input.docx> <output.txt>")
        sys.exit(1)
    
    extract_docx_text(sys.argv[1], sys.argv[2])
    print(f"Extracted {sys.argv[1]} to {sys.argv[2]}")
